"""Reconcile the workbook, the dashboard and the broker - and say where they differ.

Run this after every upload. It answers the question a committee member asks when
they have the Excel open on one screen and the dashboard on the other: are these
the same numbers, and which one do I trust?

    python tools/reconcile.py "C:\\path\\CEE Fund Portfolios - Master.xlsx"

Everything is read fresh. Nothing is written, nothing is published, and no part of
the dashboard's own code is imported - the point is for this to be able to
disagree with it.

What it checks, in order of how much it matters:

  1. The Holdings Snapshot against the broker's position exports. The snapshot is
     what every holdings-driven figure on the dashboard is built from, so if it
     matches the custodian, those figures are right by construction.
  2. What the dashboard publishes against that snapshot.
  3. The transaction ledger against the broker: which tickets are missing, and
     what they are worth. This is what the performance history is rebuilt from.
  4. The published NAV against the Schwab month-end values typed into the
     workbook's Statement Check tab.
  5. The end of the published series against today's holdings.
  6. The workbook against itself - stale cached formulas, sold-out rows, and the
     summary cells that disagree with the rows they total.

Broker exports are picked up from the Downloads folder by filename: the newest
"...Positions-YYYY-MM-DD-...csv" for each fund, and the one before it.
"""
import sys, os, re, csv, json, glob, collections, datetime, urllib.request

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is needed: python -m pip install openpyxl")

DB = "https://cee-fund-dashboard-640ab-default-rtdb.firebaseio.com"
DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")
FUNDS = ("Endowment", "CEE Fund")
KEY = {"Endowment": "endowment", "CEE Fund": "ceeFund"}
# the export filenames the custodian produces for each account
EXPORT_GLOB = {"CEE Fund": "CEE Fund Portfolio-Positions-*.csv",
               "Endowment": "CEE Fund Mngd Endow-Positions-*.csv"}
# an old symbol that is the same security as a current one
ALIAS = {"BRK/B": "BRK.B", "30231G102": "XOM"}
SH_IN = {"Buy", "Reinvest Shares", "Spin-off", "Reinvestment Adj"}
SH_OUT = {"Sell"}
SH_DELTA = {"Mandatory Reorg Exc", "Cash Merger Adj", "Stock Split"}

findings = []
note = lambda sev, text: findings.append((sev, text))
CUSIP = re.compile(r"^[0-9][0-9A-Z]{8}$")
TICKERISH = re.compile(r"^[A-Z][A-Z0-9.\- ]{0,9}$")


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        t = v.replace("$", "").replace(",", "").strip()
        if t.startswith("(") and t.endswith(")"):
            t = "-" + t[1:-1]
        try:
            return float(t)
        except ValueError:
            return None
    return None


def iso(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "")[:10]
    return s if re.match(r"^\d{4}-\d{2}-\d{2}$", s) else None


def tk_of(v):
    t = re.sub(r"\s+", ".", str(v or "").strip().upper())
    return ALIAS.get(t, t)


def header_row(rows, *required):
    for i, r in enumerate(rows[:80]):
        labels = [str(c or "").strip() for c in (r or [])]
        if all(any(lab == want for lab in labels) for want in required):
            cols = {}
            for j, lab in enumerate(labels):
                cols.setdefault(lab, j)       # a right-hand block repeats the labels
            return i, cols
    return None, None


# ----------------------------------------------------------------- the workbook
def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {n: list(wb[n].iter_rows(values_only=True)) for n in wb.sheetnames}
    wbv = openpyxl.load_workbook(path, data_only=True)

    rows = sheets["Holdings Snapshot"]
    hi, hc = header_row(rows, "Fund", "Ticker (tab)", "Shares")
    snap, cash, asof = collections.defaultdict(dict), {}, None
    for r in rows:
        if r and str(r[0] or "").strip() == "Snapshot date":
            asof = iso(r[1])
    for r in rows[hi + 1:]:
        if not r:
            continue
        fund = str(r[0] or "").strip()
        sym = r[hc["Price symbol"]] if hc.get("Price symbol") is not None and len(r) > hc["Price symbol"] and r[hc["Price symbol"]] else r[1]
        sh = num(r[hc["Shares"]])
        if fund in FUNDS and sh is not None and abs(sh) > 1e-9:
            snap[fund][tk_of(sym)] = dict(shares=sh, cost=num(r[hc["Cost basis"]]) or 0.0)
        f2 = str(r[6] or "").strip() if len(r) > 6 else ""
        c2 = num(r[7]) if len(r) > 7 else None
        if f2 in FUNDS and c2 is not None:
            cash[f2] = c2

    tabs = {}
    for name in FUNDS:
        rws = sheets[name]
        hi2, hc2 = header_row(rws, "Ticker", "Shares", "Market Value")
        out = dict(holdings={}, bonds={}, zero={}, cash=None, stale=[])
        kind = "equity"
        get = lambda r, k: (r[hc2[k]] if hc2.get(k) is not None and len(r) > hc2[k] else None)
        for r in rws[hi2 + 1:]:
            if not r:
                continue
            label = str(get(r, "Ticker") or "").strip()
            shares = num(get(r, "Shares"))
            if label.lower().startswith("cash"):
                out["cash"] = num(get(r, "Company")) or num(get(r, "Shares"))
                continue
            if label and shares is None:
                low = label.lower()
                kind = "etf" if "etf" in low else "bond" if re.search(r"treasur|bond|fixed", low) else "equity"
                continue
            tk = tk_of(label)
            if not (TICKERISH.match(tk) or CUSIP.match(tk)) or shares is None:
                continue
            if CUSIP.match(tk):
                out["bonds"][tk] = dict(shares=shares, cost=num(get(r, "Cost Basis")),
                                        mv=num(r[6]) if len(r) > 6 else None)
                continue
            rec = dict(shares=shares, cost=num(get(r, "Cost Basis")), price=num(get(r, "Market Price")),
                       mv=num(get(r, "Market Value")), sector=str(get(r, "Sector") or "").strip(), kind=kind)
            if abs(shares) < 1e-9:
                out["zero"][tk] = rec
                continue
            if rec["price"] and rec["mv"] is not None:
                want = shares * rec["price"]
                if abs(want - rec["mv"]) > max(0.02, abs(want) * 1e-6):
                    out["stale"].append((tk, rec["mv"], want))
            out["holdings"][tk] = rec
        tabs[name] = out

    rws = sheets["Transactions"]
    ti, tc = header_row(rws, "Date", "Fund", "Action")
    tx = []
    for r in rws[ti + 1:]:
        if not r:
            continue
        cell = lambda k: (r[tc[k]] if tc.get(k) is not None and len(r) > tc[k] else None)
        d = iso(cell("Date"))
        if not d:
            continue
        tx.append(dict(date=d, fund=str(cell("Fund") or "").strip(), ticker=tk_of(cell("Ticker")),
                       action=str(cell("Action") or "").strip(), shares=num(cell("Shares")) or 0.0,
                       price=num(cell("Price")) or 0.0, amount=num(cell("Amount")) or 0.0,
                       fees=num(cell("Fees")) or 0.0))
    tx.sort(key=lambda t: t["date"])

    # month-end account values typed in from the statements
    stmt = {}
    if "Statement Check" in sheets:
        rws = sheets["Statement Check"]
        si, sc = header_row(rws, "Month end", "Fund", "Ending value")
        if si is not None:
            for r in rws[si + 1:]:
                if not r:
                    continue
                d, f = iso(r[sc["Month end"]]), str(r[sc["Fund"]] or "").strip()
                v = num(r[sc["Ending value"]])
                if d and f in FUNDS and v:
                    stmt[(f, d)] = v
    return dict(snap=snap, cash=cash, asof=asof, tabs=tabs, tx=tx, stmt=stmt, cells=wbv, sheets=sheets)


# -------------------------------------------------------------------- the broker
def read_export(path):
    out, cash, asof = {}, None, None
    rows = list(csv.reader(open(path, newline="", encoding="utf-8-sig")))
    for r in rows[:3]:
        if r and r[0].startswith("Positions for account"):
            m = re.search(r"as of ([^,]+, \d{4}/\d{2}/\d{2})", r[0])
            asof = m.group(1) if m else None
    hi = next((j for j, r in enumerate(rows) if r and r[0] == "Symbol"), None)
    if hi is None:
        return None
    cols = {c: j for j, c in enumerate(rows[hi])}
    cash_of = lambda r: num(re.sub(r"[$,]", "", r[cols["Mkt Val (Market Value)"]]))
    for r in rows[hi + 1:]:
        if not r or not r[0]:
            continue
        if r[0].lower().startswith("cash"):
            cash = cash_of(r)
            continue
        q = num(str(r[cols["Qty (Quantity)"]]).replace(",", ""))
        if q is None:
            continue
        out[tk_of(r[0])] = dict(shares=q, cost=num(re.sub(r"[$,]", "", r[cols["Cost Basis"]])) or 0.0,
                                price=num(str(r[cols["Price"]]).replace(",", "")) or 0.0, mv=cash_of(r) or 0.0)
    return dict(pos=out, cash=cash, asof=asof, path=path)


def latest_exports(fund):
    files = sorted(glob.glob(os.path.join(DOWNLOADS, EXPORT_GLOB[fund])))
    dated = []
    for f in files:
        m = re.search(r"(\d{4}-\d{2}-\d{2})", os.path.basename(f))
        if m:
            dated.append((m.group(1), f))
    dated.sort()
    return [read_export(f) for _, f in dated[-2:]][::-1] if dated else []


def fetch(node):
    with urllib.request.urlopen(f"{DB}/{node}.json", timeout=240) as r:
        return json.load(r)


# ------------------------------------------------------------------------ report
def main():
    path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DOWNLOADS, "CEE Fund Portfolios - Master.xlsx")
    if not os.path.exists(path):
        sys.exit(f"workbook not found: {path}")
    W = read_workbook(path)
    print("=" * 78)
    print("RECONCILIATION")
    print(f"workbook   {os.path.basename(path)}  (snapshot {W['asof']}, {len(W['tx'])} tickets)")

    H = fetch("ceeHoldings")
    P = fetch("ceePerformance")
    PH = fetch("ceePerformanceHoldings")
    print(f"dashboard  {H.get('source')}  published {str(H.get('updatedAt'))[:16]}")
    print(f"series     {len(P['dates'])} days to {P['dates'][-1]}, anchor {P.get('anchor')}")
    print("=" * 78)

    # ---- 1. snapshot vs broker
    print("\n[1] Holdings Snapshot vs the broker's own position export")
    exports = {}
    for fund in FUNDS:
        got = latest_exports(fund)
        if not got or not got[0]:
            note("WARN", f"no broker export found for {fund} in {DOWNLOADS}")
            continue
        exports[fund] = got
        b = got[0]
        s = dict(W["snap"][fund])
        for tk, v in W["tabs"][fund]["bonds"].items():
            s[tk] = v
        bad = 0
        for tk in sorted(set(b["pos"]) | set(s)):
            bb, ss = b["pos"].get(tk), s.get(tk)
            if not bb or not ss:
                note("HIGH", f"{fund} {tk}: " + ("at the broker, not in the workbook" if bb else "in the workbook, not at the broker"))
                bad += 1
                continue
            if abs(bb["shares"] - ss["shares"]) > 1e-4:
                note("HIGH", f"{fund} {tk}: workbook {ss['shares']:g} sh vs broker {bb['shares']:g}")
                bad += 1
            if abs(bb["cost"] - (ss.get("cost") or 0)) > max(1.0, bb["cost"] * 0.005):
                note("MED", f"{fund} {tk}: cost basis ${ss.get('cost'):,.2f} vs broker ${bb['cost']:,.2f}")
                bad += 1
        wc, bc = W["tabs"][fund]["cash"] or 0, b["cash"] or 0
        if abs(wc - bc) > 0.02:
            note("HIGH", f"{fund}: cash ${wc:,.2f} vs broker ${bc:,.2f}")
            bad += 1
        print(f"   {fund}: {len(b['pos'])} positions as of {b['asof']} - "
              f"{'matches to the share and the cent' if not bad else f'{bad} difference(s)'}")

    # ---- 2. dashboard vs snapshot
    print("\n[2] What the dashboard publishes vs the snapshot")
    for fund in FUNDS:
        pub = {}
        for h in (H[KEY[fund]].get("equities") or []) + (H[KEY[fund]].get("etfs") or []):
            pub[tk_of(h["ticker"])] = h
        bad = 0
        for tk in sorted(set(pub) | set(W["snap"][fund])):
            g, s = pub.get(tk), W["snap"][fund].get(tk)
            if not g or not s:
                note("HIGH", f"{fund} {tk}: " + ("published but not in the snapshot" if g else "in the snapshot but not published"))
                bad += 1
                continue
            if abs((g.get("shares") or 0) - s["shares"]) > 1e-4 or abs((g.get("costBasis") or 0) - s["cost"]) > 0.02:
                note("HIGH", f"{fund} {tk}: published {g.get('shares')} sh / ${g.get('costBasis'):,.2f} "
                             f"vs snapshot {s['shares']:g} sh / ${s['cost']:,.2f}")
                bad += 1
        pc = H[KEY[fund]].get("cash") or 0
        if abs(pc - (W["cash"].get(fund) or 0)) > 0.02:
            note("HIGH", f"{fund}: published cash ${pc:,.2f} vs snapshot ${W['cash'].get(fund):,.2f}")
        print(f"   {fund}: {len(pub)} holdings - {'identical' if not bad else f'{bad} difference(s)'}")
    if len(H.get("transactions") or []) != len(W["tx"]):
        note("HIGH", f"ledger: {len(H.get('transactions') or [])} tickets published vs {len(W['tx'])} in the workbook")

    # ---- 3. the ledger against the broker
    print("\n[3] Tickets the ledger is missing (broker share changes it cannot explain)")
    # The workbook publishes its own estimates for trades the broker's transaction
    # export has not reached yet, and the rebuild folds those in, so a gap one of
    # them covers is already accounted for in the published series.
    est = collections.defaultdict(float)
    for e in (H.get("estimatedTrades") or []):
        est[(str(e.get("fund") or "").strip(), tk_of(e.get("ticker")))] += float(e.get("shares") or 0)
    if est:
        print(f"   ({len(H.get('estimatedTrades') or [])} estimated trades published with the workbook are folded in)")
    covers = lambda fund, tk, gap: abs(est.get((fund, tk), 0.0) - gap) <= max(0.01, abs(gap) * 0.05)
    for fund in FUNDS:
        got = exports.get(fund)
        if not got or len(got) < 2 or not got[1]:
            print(f"   {fund}: needs two broker exports to compare")
            continue
        new, old = got[0], got[1]
        since = re.search(r"(\d{4}-\d{2}-\d{2})", os.path.basename(old["path"])).group(1)
        led = collections.defaultdict(float)
        for t in W["tx"]:
            if t["fund"] != fund or t["date"] <= since or not t["ticker"]:
                continue
            d = t["shares"] if t["action"] in (SH_IN | SH_DELTA) else (-t["shares"] if t["action"] in SH_OUT else 0.0)
            if d:
                led[t["ticker"]] += d
        gaps = []
        for tk in set(new["pos"]) | set(old["pos"]):
            moved = new["pos"].get(tk, {}).get("shares", 0.0) - old["pos"].get(tk, {}).get("shares", 0.0)
            gap = moved - led.get(tk, 0.0)
            if abs(gap) > 0.004:
                px = (new["pos"].get(tk) or old["pos"].get(tk))["price"]
                gaps.append((tk, gap, gap * px))
        print(f"   {fund}: {since} -> {re.search(r'(\d{4}-\d{2}-\d{2})', os.path.basename(new['path'])).group(1)}"
              f" - {len(gaps)} holding(s) the ledger cannot account for")
        ncov = 0
        for tk, gap, val in sorted(gaps, key=lambda x: -abs(x[2]))[:14]:
            # a reinvested dividend is small in money terms however many shares it
            # buys - 1.38 shares of a $39 ETF is a $54 dividend, not a trade
            trade = abs(val) >= 400
            ok = covers(fund, tk, gap)
            ncov += ok
            print(f"      {tk:8} {gap:+10.4f} sh  (~${abs(val):,.0f})  "
                  f"missing {'trade' if trade else 'dividend reinvestment'}"
                  f"{'  - covered by the workbook estimate' if ok else ''}")
            if trade and not ok:
                note("HIGH", f"{fund} {tk}: neither a ticket nor an estimate for {gap:+.4f} shares "
                             f"(~${abs(val):,.0f}) - the rebuild treats the position as though it had "
                             f"always been that size")
        frac = [g for g in gaps if abs(g[2]) < 400 and not covers(fund, g[0], g[1])]
        if frac:
            note("MED", f"{fund}: {len(frac)} dividend reinvestment(s) with neither a ticket nor an "
                        f"estimate, about ${sum(abs(g[2]) for g in frac):,.0f} of shares")
        if ncov:
            note("NOTE", f"{fund}: {ncov} share change(s) the ledger lacks are covered by the workbook's "
                         f"estimated trades - swap them for the real tickets when the export catches up")

    # ---- 4. published NAV vs the statements
    print("\n[4] Published NAV against the Schwab statement month-ends")
    if not W["stmt"]:
        print("   nothing typed into the Statement Check tab yet")
    for (fund, day), stated in sorted(W["stmt"].items(), key=lambda x: (x[0][1], x[0][0])):
        idx = [i for i, d in enumerate(P["dates"]) if d <= day]
        if not idx:
            continue
        nav = P["funds"][KEY[fund]]["nav"][max(idx)]
        pct = (nav / stated - 1) * 100
        print(f"   {fund:10} {day}: dashboard ${nav:,.2f} vs statement ${stated:,.2f}  ({pct:+.3f}%)")
        if abs(pct) > 0.5:
            note("HIGH", f"{fund} {day}: NAV is {pct:+.2f}% from the statement")
        elif abs(pct) > 0.1:
            note("MED", f"{fund} {day}: NAV is {pct:+.2f}% from the statement")

    # ---- 5. end of the series vs today's holdings
    print("\n[5] End of the series vs the holdings it should land on")
    px = {r[0]: (r[1], r[2]) for r in PH["px"]}
    i = len(P["dates"]) - 1
    for fund in FUNDS:
        eng = {}
        for row in PH["funds"][KEY[fund]]:
            for start, vals in row[1]:
                if start <= i < start + len(vals):
                    eng[row[0]] = vals[i - start]
        want, diffs = 0.0, []
        held = dict(W["snap"][fund])
        for tk, v in W["tabs"][fund]["bonds"].items():
            held[tk] = v
        for tk, rec in held.items():
            pr = None
            if tk in px and 0 <= i - px[tk][0] < len(px[tk][1]):
                pr = px[tk][1][i - px[tk][0]]
            elif CUSIP.match(tk) and rec.get("mv") and rec.get("shares"):
                pr = rec["mv"] / rec["shares"]
            if pr is None:
                continue
            w = rec["shares"] * pr
            want += w
            if abs(eng.get(tk, 0.0) - w) > max(1.0, abs(w) * 0.002):
                diffs.append((tk, rec["shares"], eng.get(tk, 0.0) / pr if pr else 0, eng.get(tk, 0.0) - w))
        nav = P["funds"][KEY[fund]]["nav"][i]
        target = want + (W["cash"].get(fund) or 0)
        print(f"   {fund}: series ${nav:,.2f} vs holdings x {P['dates'][i]} close + cash ${target:,.2f} "
              f"({nav - target:+,.2f})")
        for tk, sh, esh, d in sorted(diffs, key=lambda x: -abs(x[3]))[:6]:
            note("HIGH" if abs(d) > 500 else "MED",
                 f"{fund} {tk}: series carries {esh:,.4f} sh, the fund holds {sh:,.4f} (${d:+,.2f})")

    # ---- 6. the workbook against itself
    print("\n[6] The workbook against itself")
    c = W["cells"]
    totals = {"Endowment rows (Q16)": c["Endowment"]["Q16"].value, "CEE Fund rows (P16)": c["CEE Fund"]["P16"].value,
              "Dashboard tab Endowment": c["Dashboard"]["C13"].value, "Dashboard tab CEE Fund": c["Dashboard"]["D13"].value,
              "Dashboard tab total": c["Dashboard"]["E13"].value, "Total tab": c["Total"]["B1"].value}
    for k, v in totals.items():
        if isinstance(v, (int, float)):
            print(f"   {k:28} ${v:,.2f}")
    a, b = totals["Dashboard tab total"], totals["Total tab"]
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) > 1:
        note("MED", f"the workbook reports two different totals: Dashboard tab ${a:,.2f} vs Total tab ${b:,.2f} "
                    f"- cached results from different moments; opening the file recalculates them")
    for fund in FUNDS:
        t = W["tabs"][fund]
        if t["stale"]:
            worst = max(t["stale"], key=lambda x: abs(x[1] - x[2]))
            note("MED", f"{fund}: {len(t['stale'])} row(s) whose saved Market Value is not Shares x Price, "
                        f"worst {worst[0]} ${worst[1]:,.2f} -> ${worst[2]:,.2f} - the file was saved before recalculating")
        if t["zero"]:
            tot = sum(z["mv"] or 0 for z in t["zero"].values())
            note("MED", f"{fund}: {len(t['zero'])} sold-out row(s) still on the tab carrying ${tot:,.2f} "
                        f"({', '.join(t['zero'])}) - they are counted in the tab's position count")
        shown = c["Dashboard"]["C19" if fund == "Endowment" else "D19"].value
        if isinstance(shown, (int, float)) and shown != len(t["holdings"]):
            note("MED", f"{fund}: the workbook counts {shown:g} positions, {len(t['holdings'])} have shares")

    print("\n" + "=" * 78)
    print("FINDINGS")
    print("=" * 78)
    rank = {"HIGH": 0, "MED": 1, "WARN": 2, "NOTE": 3}
    for sev, text in sorted(findings, key=lambda x: rank.get(x[0], 3)):
        print(f"  [{sev:4}] {text}")
    if not findings:
        print("  nothing to report - the workbook, the dashboard and the broker agree")
    print(f"\n{len(findings)} finding(s)")
    return 1 if any(s == "HIGH" for s, _ in findings) else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
