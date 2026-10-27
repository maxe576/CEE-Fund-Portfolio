"""Work out the S&P 500's current sector weights and publish them.

The dashboard reports where each fund is overweight or underweight against the
index, and that comparison is only as good as the index weights behind it. They
used to be a table typed into the page, which went stale: information technology
had moved from roughly 32 per cent to nearly 40 while the page still said 32.

Two sources, both plain and both free:

  * State Street publishes SPY's holdings every day, as an Excel file, with each
    holding's exact index weight and the date it is good for. That is the
    authority on weights - it is the fund tracking the index - but its sector
    column is left blank.
  * Wikipedia's list of S&P 500 companies carries the GICS sector for every
    constituent, and is kept current within a day or two of a change.

Joining the two gives each GICS sector's weight. The weights are the index's own;
only the sector labels come from Wikipedia, and the script refuses to publish if
it cannot label enough of the index by weight.

    python tools/index-weights.py            # fetch, check, publish
    python tools/index-weights.py --dry-run  # fetch and print, publish nothing

Run by the nightly workflow. Writes one small Firebase node, /ceeIndexSectors,
and touches nothing else.
"""
import sys, io, json, re, html, datetime, urllib.request, urllib.error

SPY_HOLDINGS = ("https://www.ssga.com/library-content/products/fund-data/etfs/us/"
                "holdings-daily-us-en-spy.xlsx")
WIKI = "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies"
NODE = "https://cee-fund-dashboard-640ab-default-rtdb.firebaseio.com/ceeIndexSectors.json"
UA = {"User-Agent": "Mozilla/5.0 (CEE Fund dashboard; sector weights)"}

# The eleven GICS sectors, named as the dashboard and the workbook name them.
SECTORS = ["Information Technology", "Financials", "Health Care", "Consumer Discretionary",
           "Communication Services", "Industrials", "Consumer Staples", "Energy",
           "Utilities", "Real Estate", "Materials"]
# Wikipedia's spellings that differ from ours.
ALIAS = {"Information technology": "Information Technology", "Health care": "Health Care",
         "Consumer discretionary": "Consumer Discretionary", "Consumer staples": "Consumer Staples",
         "Communication services": "Communication Services", "Real estate": "Real Estate"}
# A ticker written differently by the two sources.
TICKER_FIX = {"BRK.B": "BRK.B", "BF.B": "BF.B", "BRK/B": "BRK.B", "BF/B": "BF.B"}

MIN_COVERAGE = 97.0        # per cent of index weight that must carry a sector


def fetch(url, binary=False):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=120) as r:
        raw = r.read()
    return raw if binary else raw.decode("utf-8", "replace")


def spy_weights():
    """Ticker to index weight, and the date State Street says it is good for."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is needed: python -m pip install openpyxl")
    data = fetch(SPY_HOLDINGS, binary=True)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    asof = None
    for r in rows[:6]:
        for c in r:
            m = re.search(r"As of (\d{1,2}-\w{3}-\d{4})", str(c or ""))
            if m:
                asof = datetime.datetime.strptime(m.group(1), "%d-%b-%Y").date().isoformat()
    hdr = next(i for i, r in enumerate(rows)
               if r and "Ticker" in [str(c or "").strip() for c in r]
               and "Weight" in [str(c or "").strip() for c in r])
    cols = {}
    for j, c in enumerate(rows[hdr]):
        cols.setdefault(str(c or "").strip(), j)
    out = {}
    for r in rows[hdr + 1:]:
        if not r:
            continue
        tk = str(r[cols["Ticker"]] or "").strip().upper()
        try:
            w = float(str(r[cols["Weight"]]).replace(",", ""))
        except (TypeError, ValueError):
            continue
        if not tk or tk == "-" or w <= 0:
            continue
        tk = TICKER_FIX.get(tk, tk)
        out[tk] = out.get(tk, 0.0) + w      # two share classes, one line each
    return out, asof


def wiki_sectors():
    """Ticker to GICS sector, from the constituents table."""
    page = fetch(WIKI)
    start = page.find('id="constituents"')
    table = page[start:page.find("</table>", start)] if start > 0 else page
    out = {}
    # Wikipedia's rows and cells carry attributes, so the tags cannot be matched bare.
    for row in re.findall(r"<tr[^>]*>(.*?)</tr>", table, re.S):
        cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, re.S)
        if len(cells) < 4:
            continue
        strip = lambda s: html.unescape(re.sub(r"<[^>]+>", "", s)).strip()
        tk, sector = strip(cells[0]).upper().replace("​", ""), strip(cells[2])
        sector = ALIAS.get(sector, sector)
        if re.match(r"^[A-Z][A-Z.\-]{0,8}$", tk) and sector in SECTORS:
            out[TICKER_FIX.get(tk, tk)] = sector
    return out


def main():
    dry = "--dry-run" in sys.argv
    weights, asof = spy_weights()
    sectors = wiki_sectors()
    print(f"SPY holdings: {len(weights)} tickers, {sum(weights.values()):.2f}% total weight, as of {asof}")
    print(f"GICS sectors: {len(sectors)} tickers labelled")

    totals = {s: 0.0 for s in SECTORS}
    unlabelled = []
    for tk, w in weights.items():
        s = sectors.get(tk)
        if s:
            totals[s] += w
        else:
            unlabelled.append((tk, w))
    labelled = sum(totals.values())
    missing = sum(w for _, w in unlabelled)
    coverage = labelled / (labelled + missing) * 100 if (labelled + missing) else 0

    print(f"labelled {labelled:.2f}% of weight, {missing:.2f}% not labelled "
          f"({len(unlabelled)} tickers) - coverage {coverage:.2f}%")
    for tk, w in sorted(unlabelled, key=lambda x: -x[1])[:10]:
        print(f"    unlabelled: {tk} {w:.3f}%")

    # Normalise to 100 so the comparison against fund weights is like for like:
    # the fund's sector weights are a share of its equity and ETF value, and the
    # index's residual here is cash and the few names we could not label.
    out = {s: round(totals[s] / labelled * 100, 2) for s in SECTORS}
    print("\n  sector weights (normalised to 100):")
    for s, w in sorted(out.items(), key=lambda x: -x[1]):
        print(f"    {s:24} {w:6.2f}%")
    print(f"    {'sum':24} {sum(out.values()):6.2f}%")

    if coverage < MIN_COVERAGE:
        sys.exit(f"\nNOT PUBLISHED - only {coverage:.2f}% of the index could be labelled, "
                 f"below the {MIN_COVERAGE}% floor. A sector table that quietly drops part of "
                 f"the index would make every overweight look larger than it is.")
    if any(v <= 0 for v in out.values()):
        sys.exit("\nNOT PUBLISHED - a sector came out at zero, which means the join failed.")

    payload = {"asOf": asof, "updatedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
               "weights": out, "coverage": round(coverage, 2),
               "holdings": len(weights),
               "source": "SPDR S&P 500 ETF daily holdings (weights) joined with the GICS sector "
                         "of each constituent"}
    if dry:
        print("\n--dry-run: nothing published")
        return 0
    req = urllib.request.Request(NODE, data=json.dumps(payload).encode(),
                                 method="PUT", headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as r:
        print(f"\npublished to /ceeIndexSectors: {r.status}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.exit(main())
